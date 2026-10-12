"""End-to-end tests of the platform admin dashboard against a real PostgreSQL database.

The dashboard is driven through Flask's test client with a fake authenticator, and talks to
the database as the real ``saas_platform_admin`` and ``app_backend`` logins. Tenants created
here are committed (the app uses its own connections), so every test cleans them up.

The tenants onboarded in these tests are deliberately unrelated businesses: the same
generic code path must serve them all.
"""

from __future__ import annotations

import io
import json
import re
import unittest
import uuid

from openpyxl import load_workbook

from platform_admin.app import create_app
from platform_admin.auth import AuthError, Authenticator, Identity
from platform_admin.config import Settings
from platform_admin.services.catalogue_import import template_bytes
from saas.bot import BotProviderError, BotResponseRequest
from saas_test_support import (
    PLATFORM_ADMIN_USER_ID,
    SaasDatabaseTestCase,
    _connect,
    app_backend_test_url,
    platform_admin_test_url,
    TENANT_ONE_OWNER_USER_ID,
)

SLUG_PREFIX = "pa-test-"
PASSWORD = "correct horse"


class FakeBotProvider:
    def __init__(self):
        self.requests: list[BotResponseRequest] = []
        self.failure: Exception | None = None

    def generate(self, request: BotResponseRequest) -> str:
        self.requests.append(request)
        if self.failure is not None:
            raise self.failure
        return f"Grounded AI response: {request.customer_message}"


class FakeAuthenticator(Authenticator):
    def __init__(self, users: dict[str, tuple[str, str]]):
        self._users = users

    def login(self, email, password):
        entry = self._users.get(email)
        if not entry or entry[0] != password:
            raise AuthError("Invalid email or password.")
        return Identity(entry[1], email)


SPECS = (
    {
        "name": "Alpha Salon", "slug": SLUG_PREFIX + "alpha", "currency": "KES",
        "rules": {"deposit_rate": 0.5, "services": ["cut", "colour"]},
        "persona": {"tone": "friendly"},
        "hours": {"mon": "09:00-17:00"},
        "product": {"name": "Haircut", "category": "service"},
        "variant": {"code": "std", "name": "Standard", "price": "1500.00"},
    },
    {
        "name": "Beta Repairs", "slug": SLUG_PREFIX + "beta", "currency": "USD",
        "rules": {"call_out_fee": 20, "warranty_days": 90},
        "persona": {"tone": "formal"},
        "hours": {"sat": "10:00-14:00"},
        "product": {"name": "Screen repair", "category": "repair"},
        "variant": {"code": "oem", "name": "OEM part", "price": "89.50"},
    },
    {
        "name": "Cindy Bakes", "slug": SLUG_PREFIX + "bakes", "currency": "KES",
        "rules": {"deposit_rate": 0.7, "delivery_rate_per_km": 32},
        "persona": {"tone": "warm"},
        "hours": {"mon-sat": "08:00-18:00"},
        "product": {"name": "Vanilla cake", "category": "cake"},
        "variant": {"code": "1kg", "name": "1kg", "price": "2500.00"},
    },
)


def _settings(**overrides) -> Settings:
    return Settings(
        secret_key="k" * 40,
        cookie_secure=False,
        app_database_url=app_backend_test_url(),
        platform_database_url=platform_admin_test_url(),
        **overrides,
    )


class DashboardTestCase(SaasDatabaseTestCase):
    def setUp(self):
        self._cleanup()
        conn = _connect()
        with conn, conn.cursor() as cursor:
            cursor.execute(
                "insert into public.platform_admins (user_id) values (%s) on conflict do nothing",
                (PLATFORM_ADMIN_USER_ID,),
            )
        conn.close()
        self.authenticator = FakeAuthenticator({
            "admin@example.com": (PASSWORD, PLATFORM_ADMIN_USER_ID),
            "owner@example.com": (PASSWORD, TENANT_ONE_OWNER_USER_ID),
        })
        self.bot_provider = FakeBotProvider()
        self.app = create_app(
            _settings(), authenticator=self.authenticator, bot_provider=self.bot_provider
        )
        self.client = self.app.test_client()
        self.addCleanup(self._cleanup)

    @staticmethod
    def _cleanup():
        conn = _connect()
        with conn, conn.cursor() as cursor:
            cursor.execute("select id from public.tenants where slug like %s", (SLUG_PREFIX + "%",))
            ids = [row[0] for row in cursor.fetchall()]
            if ids:
                cursor.execute("delete from public.audit_log where tenant_id = any(%s)", (ids,))
                cursor.execute("delete from public.tenants where id = any(%s)", (ids,))
            cursor.execute("delete from public.reserved_phone_numbers where phone_number_id like 'pa-reserved-%'")
            cursor.execute(
                "delete from public.platform_admins where user_id = %s", (PLATFORM_ADMIN_USER_ID,)
            )
        conn.close()

    @staticmethod
    def _delete_exact_test_tenant(tenant_id: str, slug: str) -> None:
        conn = _connect()
        with conn, conn.cursor() as cursor:
            cursor.execute(
                "delete from public.audit_log where tenant_id in "
                "(select id from public.tenants where id = %s and slug = %s)",
                (tenant_id, slug),
            )
            cursor.execute(
                "delete from public.tenants where id = %s and slug = %s",
                (tenant_id, slug),
            )
        conn.close()

    # -- helpers -----------------------------------------------------------------
    @staticmethod
    def token(client) -> str:
        with client.session_transaction() as session:
            return session["csrf_token"]

    def login(self, client=None, email="admin@example.com", password=PASSWORD):
        client = client or self.client
        client.get("/login")
        return client.post(
            "/login", data={"csrf_token": self.token(client), "email": email, "password": password}
        )

    def post(self, path, data=None, client=None):
        client = client or self.client
        payload = dict(data or {})
        payload["csrf_token"] = self.token(client)
        return client.post(path, data=payload)

    def message_after(self, response, client=None) -> str:
        """The page a redirect leads to (where flash messages are shown)."""
        return (client or self.client).get(response.headers["Location"]).get_data(as_text=True)

    def create_tenant(self, spec, **extra):
        response = self.post("/tenants/new", {
            "name": spec["name"], "slug": spec["slug"], "timezone": "Africa/Nairobi",
            "currency": spec["currency"], **extra,
        })
        self.assertEqual(response.status_code, 302, response.get_data(as_text=True))
        return re.search(r"/tenants/([0-9a-f-]{36})", response.headers["Location"]).group(1)

    def onboard(self, spec) -> str:
        """The one generic workflow: create -> profile -> rules -> catalogue -> activate."""
        tenant_id = self.create_tenant(spec, owner_user_id=str(uuid.uuid4()))
        base = f"/tenants/{tenant_id}"
        self.assertEqual(self.post(f"{base}/profile", {
            "name": spec["name"], "timezone": "Africa/Nairobi", "currency": spec["currency"]}).status_code, 302)
        self.assertEqual(self.post(f"{base}/rules", {
            "business_rules": json.dumps(spec["rules"]), "bot_persona": json.dumps(spec["persona"]),
            "business_hours": json.dumps(spec["hours"]), "configuration": "{}"}).status_code, 302)
        self.post(f"{base}/catalogue/products", {**spec["product"], "sort_order": "1"})
        product_id = self.product_id(tenant_id)
        self.post(f"{base}/catalogue/products/{product_id}/variants", spec["variant"])
        self.post(f"{base}/status", {"status": "active"})
        return tenant_id

    def product_id(self, tenant_id) -> str:
        with self.superuser() as cursor:
            cursor.execute("select id::text from public.products where tenant_id = %s", (tenant_id,))
            return cursor.fetchone()[0]

    def tenant_status(self, tenant_id) -> str:
        with self.superuser() as cursor:
            cursor.execute("select status from public.tenants where id = %s", (tenant_id,))
            return cursor.fetchone()[0]

    def audit_actions(self, tenant_id) -> list[str]:
        with self.superuser() as cursor:
            cursor.execute(
                "select action from public.audit_log where tenant_id = %s order by id", (tenant_id,)
            )
            return [row[0] for row in cursor.fetchall()]

    def page(self, path) -> str:
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200, path)
        return response.get_data(as_text=True)


class LoginTests(DashboardTestCase):
    def test_administrator_can_sign_in_and_view_the_overview(self):
        self.assertEqual(self.login().status_code, 302)
        self.assertIn("Total tenants", self.page("/"))

    def test_wrong_password_is_rejected_without_detail(self):
        response = self.login(password="nope")
        self.assertEqual(response.status_code, 401)
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_valid_user_who_is_not_a_platform_admin_gets_the_same_generic_error(self):
        wrong = self.login(password="nope").get_data(as_text=True)
        non_admin = self.login(email="owner@example.com").get_data(as_text=True)
        self.assertIn("Invalid email or password.", non_admin)
        self.assertIn("Invalid email or password.", wrong)
        self.assertEqual(self.client.get("/tenants").status_code, 302)

    def test_repeated_failures_are_rate_limited(self):
        statuses = [self.login(password="bad").status_code for _ in range(7)]
        self.assertEqual(statuses[:5], [401] * 5)
        self.assertEqual(statuses[-1], 429)

    def test_rate_limiting_still_lets_other_addresses_sign_in(self):
        for _ in range(6):
            self.login(password="bad")
        other = self.app.test_client()
        other.get("/login")
        response = other.post("/login", data={
            "csrf_token": self.token(other), "email": "admin@example.com", "password": PASSWORD},
            environ_overrides={"REMOTE_ADDR": "198.51.100.20"})
        self.assertEqual(response.status_code, 302)

    def test_forged_session_for_a_non_administrator_is_forbidden(self):
        self.login()
        with self.client.session_transaction() as session:
            session["user_id"] = TENANT_ONE_OWNER_USER_ID
        self.assertEqual(self.client.get("/").status_code, 403)
        self.assertEqual(self.client.get("/").status_code, 302)  # session was cleared

    def test_a_removed_administrator_loses_access_immediately(self):
        self.login()
        self.assertEqual(self.client.get("/").status_code, 200)
        conn = _connect()
        with conn, conn.cursor() as cursor:
            cursor.execute("update public.platform_admins set status = 'disabled' where user_id = %s",
                           (PLATFORM_ADMIN_USER_ID,))
        conn.close()
        self.assertEqual(self.client.get("/").status_code, 403)

    def test_sign_out_ends_the_session(self):
        self.login()
        self.post("/logout")
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_the_session_cookie_never_contains_a_token(self):
        self.assertNotIn("eyJ", self.login().headers.get("Set-Cookie", ""))


class OnboardingTests(DashboardTestCase):
    def setUp(self):
        super().setUp()
        self.login()

    def test_new_tenants_start_as_drafts(self):
        tenant_id = self.create_tenant(SPECS[0])
        self.assertEqual(self.tenant_status(tenant_id), "draft")
        self.assertIn("draft", self.page(f"/tenants/{tenant_id}"))

    def test_cindy_bakes_and_unrelated_businesses_onboard_through_the_same_workflow(self):
        ids = [self.onboard(spec) for spec in SPECS]
        for tenant_id, spec in zip(ids, SPECS, strict=True):
            self.assertEqual(self.tenant_status(tenant_id), "active", spec["name"])
            rules = self.page(f"/tenants/{tenant_id}/rules")
            for key in spec["rules"]:
                self.assertIn(key, rules)
            for other in (s for s in SPECS if s is not spec):
                for key in set(other["rules"]) - set(spec["rules"]):
                    self.assertNotIn(key, rules)
                self.assertNotIn(other["product"]["name"], self.page(f"/tenants/{tenant_id}/catalogue"))
            self.assertIn(spec["product"]["name"], self.page(f"/tenants/{tenant_id}/catalogue"))
            actions = self.audit_actions(tenant_id)
            for expected in ("tenant.provisioned", "tenant.profile_updated", "tenant.rules_updated",
                             "catalogue.product_created", "catalogue.variant_created",
                             "tenant.status_changed"):
                self.assertIn(expected, actions)

    def test_settings_are_stored_per_tenant(self):
        ids = [self.onboard(spec) for spec in SPECS[:2]]
        with self.superuser() as cursor:
            cursor.execute(
                "select tenant_id::text, currency, business_rules from public.tenant_settings "
                "where tenant_id = any(%s::uuid[])", (ids,))
            stored = {row[0]: row[1:] for row in cursor.fetchall()}
        self.assertEqual(stored[ids[0]][1], SPECS[0]["rules"])
        self.assertEqual(stored[ids[1]][1], SPECS[1]["rules"])
        self.assertEqual((stored[ids[0]][0], stored[ids[1]][0]), ("KES", "USD"))

    def test_activation_is_blocked_until_the_tenant_is_ready(self):
        tenant_id = self.create_tenant(SPECS[0], owner_user_id=str(uuid.uuid4()))
        response = self.post(f"/tenants/{tenant_id}/status", {"status": "active"})
        self.assertIn("not ready to activate", self.message_after(response))
        self.assertEqual(self.tenant_status(tenant_id), "draft")

    def test_activation_without_an_owner_needs_an_explicit_waiver(self):
        tenant_id = self.create_tenant(SPECS[0])
        base = f"/tenants/{tenant_id}"
        self.post(f"{base}/catalogue/products", SPECS[0]["product"])
        self.post(f"{base}/catalogue/products/{self.product_id(tenant_id)}/variants", SPECS[0]["variant"])
        self.post(f"{base}/status", {"status": "active"})
        self.assertEqual(self.tenant_status(tenant_id), "draft")
        self.post(f"{base}/status", {"status": "active", "waive_owner": "1"})
        self.assertEqual(self.tenant_status(tenant_id), "active")

    def test_unknown_status_is_a_friendly_error_not_a_500(self):
        tenant_id = self.create_tenant(SPECS[0])
        response = self.post(f"/tenants/{tenant_id}/status", {"status": "bogus"})
        self.assertEqual(response.status_code, 302)
        self.assertIn("Status must be", self.message_after(response))

    def test_duplicate_slug_is_reported(self):
        self.create_tenant(SPECS[0])
        response = self.post("/tenants/new", {"name": "Again", "slug": SPECS[0]["slug"],
                                              "timezone": "Africa/Nairobi", "currency": "KES"})
        self.assertEqual(response.status_code, 409)
        self.assertIn("slug is already in use", response.get_data(as_text=True))

    def test_invalid_input_is_rejected_before_the_database(self):
        for data in ({"name": "", "slug": "pa-test-x"}, {"name": "X", "slug": "Bad Slug"},
                     {"name": "X", "slug": "pa-test-y", "currency": "KESS"},
                     {"name": "X", "slug": "pa-test-z", "owner_user_id": "not-a-uuid"},
                     {"name": "X\x00Y", "slug": "pa-test-nul"}):
            with self.subTest(data=data):
                response = self.post("/tenants/new", {"timezone": "Africa/Nairobi", "currency": "KES", **data})
                self.assertEqual(response.status_code, 400)

    def test_rules_must_be_json_objects(self):
        tenant_id = self.create_tenant(SPECS[0])
        response = self.post(f"/tenants/{tenant_id}/rules", {"business_rules": "[1, 2]"})
        self.assertIn("must be a JSON object", self.message_after(response))

    def test_unknown_or_malformed_tenant_ids_are_not_found(self):
        self.assertEqual(self.client.get(f"/tenants/{uuid.uuid4()}").status_code, 404)
        self.assertEqual(self.client.get("/tenants/not-a-uuid").status_code, 400)

    def test_overview_and_tenant_list_reflect_the_new_tenants(self):
        tenant_id = self.onboard(SPECS[0])
        self.assertIn(SPECS[0]["name"], self.page("/tenants"))
        self.assertIn(SPECS[0]["name"], self.page("/tenants?status=active"))
        self.assertIn("tenant.status_changed", self.page("/"))
        self.assertIn("tenant.provisioned", self.page(f"/tenants/{tenant_id}/audit"))

    def test_audit_records_the_verified_administrator(self):
        tenant_id = self.onboard(SPECS[0])
        with self.superuser() as cursor:
            cursor.execute(
                "select distinct actor_user_id::text, actor_type from public.audit_log where tenant_id = %s",
                (tenant_id,))
            self.assertEqual(cursor.fetchall(), [(PLATFORM_ADMIN_USER_ID, "platform")])

    def test_database_failures_inputs_do_not_become_500s(self):
        tenant_id = self.create_tenant(SPECS[0])
        base = f"/tenants/{tenant_id}"
        for path, data in ((f"{base}/catalogue/products", {"name": "X", "sort_order": "99999999999"}),
                           (f"{base}/rules", {"business_rules": '{"a": "\\u0000"}'}),
                           (f"{base}/catalogue/products", {"name": "N\x00"})):
            with self.subTest(path=path, data=data):
                response = self.post(path, data)
                self.assertIn(response.status_code, (302, 400))


class CatalogueTests(DashboardTestCase):
    def setUp(self):
        super().setUp()
        self.login()
        self.tenant_a = self.create_tenant(SPECS[0])
        self.tenant_b = self.create_tenant(SPECS[1])

    def add_product(self, tenant_id, name="Thing"):
        self.post(f"/tenants/{tenant_id}/catalogue/products", {"name": name})
        with self.superuser() as cursor:
            cursor.execute("select id::text from public.products where tenant_id = %s and name = %s",
                           (tenant_id, name))
            return cursor.fetchone()[0]

    @staticmethod
    def workbook(
        *,
        sku="TEST-SKU-1",
        product="Imported Item",
        variant="Default",
        price="12.50",
        currency="KES",
        availability="IN_STOCK",
    ):
        workbook = load_workbook(io.BytesIO(template_bytes()))
        sheet = workbook["Products"]
        sheet.delete_rows(3)
        values = (
            sku, product, variant, "Imported", "Imported safely", price, currency,
            availability, 5, 2, None,
        )
        for column, value in enumerate(values, 1):
            sheet.cell(2, column).value = value
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def test_products_and_variants_can_be_created_edited_priced_and_toggled(self):
        product_id = self.add_product(self.tenant_a, "Widget")
        base = f"/tenants/{self.tenant_a}/catalogue"
        self.post(f"{base}/products/{product_id}/variants", {"code": "s", "name": "Small", "price": "10"})
        with self.superuser() as cursor:
            cursor.execute("select id::text, price from public.product_variants where product_id = %s", (product_id,))
            variant_id, price = cursor.fetchone()
        self.assertEqual(str(price), "10.00")
        self.post(f"{base}/variants/{variant_id}", {"code": "s", "name": "Small", "price": "12.50"})
        self.post(f"{base}/variants/{variant_id}/active", {"active": "0"})
        self.post(f"{base}/products/{product_id}", {"name": "Widget v2", "category": "tools"})
        self.post(f"{base}/products/{product_id}/active", {"active": "0"})
        with self.superuser() as cursor:
            cursor.execute("select price, is_active from public.product_variants where id = %s", (variant_id,))
            self.assertEqual(cursor.fetchone(), (12.5, False))
            cursor.execute("select name, is_active from public.products where id = %s", (product_id,))
            self.assertEqual(cursor.fetchone(), ("Widget v2", False))
        self.assertIn("Widget v2", self.page(base))

    def test_catalogue_is_confined_to_its_tenant(self):
        product_id = self.add_product(self.tenant_a, "Private Item")
        self.assertNotIn("Private Item", self.page(f"/tenants/{self.tenant_b}/catalogue"))
        base_b = f"/tenants/{self.tenant_b}/catalogue"
        before = self.post(f"{base_b}/products/{product_id}/variants", {"code": "x", "name": "X", "price": "1"})
        self.assertIn("Product not found", self.message_after(before))
        self.post(f"{base_b}/products/{product_id}", {"name": "Hijacked"})
        self.post(f"{base_b}/products/{product_id}/active", {"active": "0"})
        with self.superuser() as cursor:
            cursor.execute("select name, is_active from public.products where id = %s", (product_id,))
            self.assertEqual(cursor.fetchone(), ("Private Item", True))
            cursor.execute("select count(*) from public.product_variants where tenant_id = %s", (self.tenant_b,))
            self.assertEqual(cursor.fetchone()[0], 0)

    def test_invalid_prices_and_duplicates_are_rejected(self):
        product_id = self.add_product(self.tenant_a)
        base = f"/tenants/{self.tenant_a}/catalogue/products/{product_id}/variants"
        for bad in ("-5", "abc", "1.999"):
            response = self.post(base, {"code": "p", "name": "P", "price": bad})
            self.assertIn("Price", self.message_after(response), bad)
        self.post(base, {"code": "dup", "name": "A", "price": "1"})
        response = self.post(base, {"code": "dup", "name": "B", "price": "2"})
        self.assertIn("already exists", self.message_after(response))

    def test_same_product_name_is_allowed_in_different_tenants(self):
        self.add_product(self.tenant_a, "Shared Name")
        self.add_product(self.tenant_b, "Shared Name")
        self.assertIn("Shared Name", self.page(f"/tenants/{self.tenant_b}/catalogue"))

    def test_excel_import_previews_then_atomically_inserts_and_updates_by_sku(self):
        base = f"/tenants/{self.tenant_a}/catalogue"
        response = self.post(
            f"{base}/import/preview",
            {"workbook": (self.workbook(), "catalogue.xlsx")},
        )
        page = response.get_data(as_text=True)
        self.assertIn("Validated changes", page)
        self.assertIn("TEST-SKU-1", page)
        batch_id = re.search(r"/import/([0-9a-f-]{36})/apply", page).group(1)
        applied = self.post(f"{base}/import/{batch_id}/apply")
        self.assertIn("Catalogue import applied", self.message_after(applied))
        with self.superuser() as cursor:
            cursor.execute(
                "select price, availability, stock_quantity from public.product_variants "
                "where tenant_id = %s and sku = 'TEST-SKU-1'",
                (self.tenant_a,),
            )
            self.assertEqual(cursor.fetchone(), (12.5, "IN_STOCK", 5))

        response = self.post(
            f"{base}/import/preview",
            {
                "workbook": (
                    self.workbook(price="15.75", availability="MADE_TO_ORDER"),
                    "catalogue.xlsx",
                )
            },
        )
        page = response.get_data(as_text=True)
        self.assertIn("availability changed", page)
        batch_id = re.search(r"/import/([0-9a-f-]{36})/apply", page).group(1)
        self.post(f"{base}/import/{batch_id}/apply")
        with self.superuser() as cursor:
            cursor.execute(
                "select count(*), min(price), min(availability) "
                "from public.product_variants where tenant_id = %s and sku = 'TEST-SKU-1'",
                (self.tenant_a,),
            )
            self.assertEqual(cursor.fetchone(), (1, 15.75, "MADE_TO_ORDER"))

    def test_import_validation_rolls_back_and_skus_are_tenant_scoped(self):
        base_a = f"/tenants/{self.tenant_a}/catalogue"
        workbook = self.workbook()
        invalid = load_workbook(workbook)
        invalid["Products"].append(
            ("BROKEN", "Broken", None, None, None, "bad", "KES", "IN_STOCK")
        )
        content = io.BytesIO()
        invalid.save(content)
        content.seek(0)
        rejected = self.post(
            f"{base_a}/import/preview",
            {"workbook": (content, "invalid.xlsx")},
        )
        self.assertIn("Import cannot be applied", rejected.get_data(as_text=True))
        with self.superuser() as cursor:
            cursor.execute(
                "select count(*) from public.product_variants where tenant_id = %s "
                "and sku in ('TEST-SKU-1', 'BROKEN')",
                (self.tenant_a,),
            )
            self.assertEqual(cursor.fetchone()[0], 0)

        for tenant_id, currency, price in (
            (self.tenant_a, "KES", "10.00"),
            (self.tenant_b, "USD", "99.00"),
        ):
            base = f"/tenants/{tenant_id}/catalogue"
            preview = self.post(
                f"{base}/import/preview",
                {
                    "workbook": (
                        self.workbook(currency=currency, price=price),
                        "catalogue.xlsx",
                    )
                },
            ).get_data(as_text=True)
            batch_id = re.search(r"/import/([0-9a-f-]{36})/apply", preview).group(1)
            self.post(f"{base}/import/{batch_id}/apply")
        with self.superuser() as cursor:
            cursor.execute(
                "select tenant_id::text, price from public.product_variants "
                "where sku = 'TEST-SKU-1' order by tenant_id",
            )
            rows = cursor.fetchall()
        self.assertEqual(len(rows), 2)
        self.assertEqual({str(price) for _, price in rows}, {"10.00", "99.00"})


class QuotationDashboardTests(DashboardTestCase):
    def setUp(self):
        super().setUp()
        self.login()
        self.tenant = self.create_tenant(SPECS[0])
        self.product_id = CatalogueTests.add_product(self, self.tenant, "Quoted Item")
        self.post(
            f"/tenants/{self.tenant}/catalogue/products/{self.product_id}/variants",
            {"sku": "QUOTE-1", "code": "q1", "name": "Standard", "price": "25.00"},
        )
        with self.superuser() as cursor:
            cursor.execute(
                "select id::text from public.product_variants where product_id = %s",
                (self.product_id,),
            )
            self.variant_id = cursor.fetchone()[0]

    def test_quotation_uses_server_prices_snapshots_values_and_downloads_pdf(self):
        response = self.post(
            f"/tenants/{self.tenant}/quotations/new",
            {
                "customer_name": "Test Customer",
                "customer_contact": "customer@example.test",
                "valid_days": "7",
                "variant_id": [self.variant_id],
                "quantity": ["2"],
                "terms": "Valid while available.",
                "price": "0.01",
            },
        )
        self.assertEqual(response.status_code, 302)
        quotation_id = re.search(
            r"/quotations/([0-9a-f-]{36})", response.headers["Location"]
        ).group(1)
        with self.superuser() as cursor:
            cursor.execute(
                "select subtotal, tax_rate, total from public.quotations where id = %s",
                (quotation_id,),
            )
            self.assertEqual(cursor.fetchone(), (50.0, None, 50.0))
            cursor.execute(
                "select sku, unit_price, quantity, availability "
                "from public.quotation_lines where quotation_id = %s",
                (quotation_id,),
            )
            self.assertEqual(cursor.fetchone(), ("QUOTE-1", 25.0, 2, "IN_STOCK"))
        self.post(
            f"/tenants/{self.tenant}/catalogue/variants/{self.variant_id}",
            {"sku": "QUOTE-1", "code": "q1", "name": "Standard", "price": "80.00"},
        )
        detail = self.page(
            f"/tenants/{self.tenant}/quotations/{quotation_id}"
        )
        self.assertIn("KES 50.00", detail)
        self.assertIn("Tax is not configured", detail)
        pdf = self.client.get(
            f"/tenants/{self.tenant}/quotations/{quotation_id}.pdf"
        )
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.data.startswith(b"%PDF-"))


class ChannelTests(DashboardTestCase):
    """Manual WhatsApp channel management. No Meta calls exist; the database enforces the rules."""

    def setUp(self):
        super().setUp()
        self.login()
        self.tenant = self.create_tenant(SPECS[0])
        self.slug = SPECS[0]["slug"]

    def connect(self, tenant_id=None, slug=None, **fields):
        data = {"phone_number_id": "pa-phone-001", "waba_id": "pa-waba-001",
                "display_phone_number": "+254700000123", "confirm_slug": slug or self.slug, **fields}
        response = self.post(f"/tenants/{tenant_id or self.tenant}/whatsapp/connect", data)
        return self.message_after(response)

    def channels(self, tenant_id=None):
        with self.superuser() as cursor:
            cursor.execute(
                "select phone_number_id, waba_id, display_phone_number, status, credentials_ref, is_default "
                "from public.tenant_channels where tenant_id = %s", (tenant_id or self.tenant,))
            return cursor.fetchall()

    def channel_id(self, tenant_id=None):
        with self.superuser() as cursor:
            cursor.execute("select id::text from public.tenant_channels where tenant_id = %s", (tenant_id or self.tenant,))
            return cursor.fetchone()[0]

    def test_a_channel_is_recorded_pending_with_no_credentials(self):
        self.connect()
        self.assertEqual(self.channels(),
                         [("pa-phone-001", "pa-waba-001", "+254700000123", "pending", None, True)])
        self.assertIn("channel.created", self.audit_actions(self.tenant))
        page = self.page(f"/tenants/{self.tenant}/whatsapp")
        self.assertIn("pa-phone-001", page)
        self.assertIn("+254700000123", page)

    def test_confirmation_slug_is_required(self):
        self.assertIn("slug exactly", self.connect(slug="wrong-slug"))
        self.assertEqual(self.channels(), [])

    def test_reserved_numbers_are_refused_by_the_database(self):
        conn = _connect()
        with conn, conn.cursor() as cursor:
            cursor.execute("insert into public.reserved_phone_numbers (phone_number_id, reason) "
                           "values ('pa-reserved-live', 'live deployment')")
        conn.close()
        self.assertIn("reserved", self.connect(phone_number_id="pa-reserved-live"))
        self.assertEqual(self.channels(), [])

    def test_invalid_identifiers_are_refused(self):
        for fields in ({"phone_number_id": "has space"}, {"waba_id": "bad id!"},
                       {"display_phone_number": "0700000000"}):
            with self.subTest(fields=fields):
                self.connect(**fields)
        self.assertEqual(self.channels(), [])

    def test_a_number_cannot_be_assigned_to_a_second_tenant(self):
        other = self.create_tenant(SPECS[1])
        self.connect()
        message = self.connect(tenant_id=other, slug=SPECS[1]["slug"])
        self.assertIn("already assigned", message)
        self.assertNotIn(SPECS[0]["name"], message)
        self.assertEqual(self.channels(other), [])

    def test_a_pending_channel_can_be_edited_and_the_change_is_audited(self):
        self.connect()
        response = self.post(f"/tenants/{self.tenant}/whatsapp/{self.channel_id()}", {
            "phone_number_id": "pa-phone-002", "waba_id": "pa-waba-002", "display_phone_number": "+254700000999"})
        self.assertIn("Channel saved", self.message_after(response))
        self.assertEqual(self.channels()[0][:3], ("pa-phone-002", "pa-waba-002", "+254700000999"))
        self.assertIn("channel.updated", self.audit_actions(self.tenant))

    def test_a_channel_of_another_tenant_cannot_be_edited_through_this_tenant(self):
        other = self.create_tenant(SPECS[1])
        self.connect(tenant_id=other, slug=SPECS[1]["slug"])
        response = self.post(f"/tenants/{self.tenant}/whatsapp/{self.channel_id(other)}", {
            "phone_number_id": "pa-phone-hijack"})
        self.assertIn("Channel not found", self.message_after(response))
        self.assertEqual(self.channels(other)[0][0], "pa-phone-001")
        response = self.post(f"/tenants/{self.tenant}/whatsapp/{self.channel_id(other)}/status", {"status": "disabled"})
        self.assertIn("Channel not found", self.message_after(response))

    def test_channels_cannot_be_activated_by_default(self):
        self.connect()
        response = self.post(f"/tenants/{self.tenant}/whatsapp/{self.channel_id()}/status", {"status": "active"})
        self.assertIn("disabled in this environment", self.message_after(response))
        self.assertEqual(self.channels()[0][3], "pending")

    def test_even_when_allowed_a_draft_tenant_cannot_have_an_active_channel(self):
        app = create_app(_settings(allow_active_channels=True), authenticator=self.authenticator)
        client = app.test_client()
        self.login(client)
        self.connect()
        response = self.post(f"/tenants/{self.tenant}/whatsapp/{self.channel_id()}/status",
                             {"status": "active"}, client=client)
        self.assertIn("requires an active tenant", self.message_after(response, client))
        self.assertEqual(self.channels()[0][3], "pending")

    def test_an_active_tenant_routes_once_its_channel_is_activated_and_identifiers_then_freeze(self):
        app = create_app(_settings(allow_active_channels=True), authenticator=self.authenticator)
        client = app.test_client()
        self.login(client)
        active_tenant = self.onboard(SPECS[1])
        self.connect(tenant_id=active_tenant, slug=SPECS[1]["slug"], phone_number_id="pa-phone-777")
        channel_id = self.channel_id(active_tenant)
        self.post(f"/tenants/{active_tenant}/whatsapp/{channel_id}/status", {"status": "active"}, client=client)
        with self.superuser() as cursor:
            cursor.execute("set local role app_backend")
            cursor.execute("select app.resolve_tenant_by_phone_number_id('pa-phone-777')")
            self.assertEqual(str(cursor.fetchone()[0]), active_tenant)
        response = self.post(f"/tenants/{active_tenant}/whatsapp/{channel_id}",
                             {"phone_number_id": "pa-phone-778"}, client=client)
        self.assertIn("disable the channel", self.message_after(response, client).lower())
        self.assertEqual(self.channels(active_tenant)[0][0], "pa-phone-777")

    def test_the_whatsapp_page_offers_no_meta_features(self):
        page = self.page(f"/tenants/{self.tenant}/whatsapp").lower()
        for word in ("discover", "list businesses", "list wabas", "check webhook", "verify with meta"):
            self.assertNotIn(word, page)
        self.assertNotRegex(page, r'<input[^>]*name="[^"]*(access|secret|password|api)')


class AuditAttributionTests(DashboardTestCase):
    def setUp(self):
        super().setUp()
        self.login()
        self.tenant = self.create_tenant(SPECS[0])

    def test_every_dashboard_audit_event_is_attributed_to_the_verified_admin_via_a_platform_function(self):
        base = f"/tenants/{self.tenant}"
        self.post(f"{base}/profile", {"name": "Alpha Salon", "timezone": "Africa/Nairobi", "currency": "KES"})
        self.post(f"{base}/rules", {"business_rules": "{}"})
        self.post(f"{base}/catalogue/products", {"name": "Thing"})
        self.post(f"{base}/members", {"user_id": str(uuid.uuid4()), "role": "staff"})
        self.post(f"{base}/whatsapp/connect", {"phone_number_id": "pa-attrib-1", "confirm_slug": SPECS[0]["slug"]})
        with self.superuser() as cursor:
            cursor.execute(
                "select distinct actor_type, actor_user_id::text from public.audit_log where tenant_id = %s",
                (self.tenant,))
            self.assertEqual(cursor.fetchall(), [("platform", PLATFORM_ADMIN_USER_ID)])
            cursor.execute("select count(*) from public.audit_log where tenant_id = %s", (self.tenant,))
            self.assertGreaterEqual(cursor.fetchone()[0], 6)

    def test_a_failed_audit_write_rolls_back_the_change(self):
        # Dropping the admin's registration makes the audit function refuse; the change must not persist.
        base = f"/tenants/{self.tenant}"
        conn = _connect()
        with conn, conn.cursor() as cursor:
            cursor.execute("update public.platform_admins set status = 'disabled' where user_id = %s",
                           (PLATFORM_ADMIN_USER_ID,))
        conn.close()
        self.post(f"{base}/catalogue/products", {"name": "Never Saved"})
        with self.superuser() as cursor:
            cursor.execute("select count(*) from public.products where tenant_id = %s", (self.tenant,))
            self.assertEqual(cursor.fetchone()[0], 0)


class BotPlaygroundTests(DashboardTestCase):
    def setUp(self):
        super().setUp()
        self.login()

    def test_playground_defaults_to_the_saas_test_tenant_when_present(self):
        with self.superuser() as cursor:
            cursor.execute("select id::text from public.tenants where slug = 'saas-test'")
            existing = cursor.fetchone()

        if existing:
            tenant_id = existing[0]
        else:
            tenant_id = self.create_tenant({
                "name": "saas-test",
                "slug": "saas-test",
                "currency": "KES",
            }, owner_user_id=TENANT_ONE_OWNER_USER_ID)
            self.addCleanup(self._delete_exact_test_tenant, tenant_id, "saas-test")
            self.post(f"/tenants/{tenant_id}/profile", {
                "name": "saas-test",
                "timezone": "Africa/Nairobi",
                "currency": "KES",
            })
            self.post(f"/tenants/{tenant_id}/rules", {
                "business_rules": json.dumps({"deposit_rate": 0.25}),
                "bot_persona": json.dumps({"tone": "friendly", "name": "SaaS Test Bot"}),
                "business_hours": json.dumps({"mon": "09:00-17:00"}),
                "configuration": json.dumps({"test_mode": True}),
            })
            self.post(
                f"/tenants/{tenant_id}/catalogue/products",
                {"name": "Lemon cake", "category": "cake"},
            )
            product_id = self.product_id(tenant_id)
            self.post(f"/tenants/{tenant_id}/catalogue/products/{product_id}/variants", {
                "code": "std", "name": "Large", "price": "500.00",
            })
            self.post(f"/tenants/{tenant_id}/status", {"status": "active"})
            self.assertEqual(self.tenant_status(tenant_id), "active")

        page = self.page("/playground")
        self.assertIn("Bot Playground", page)
        self.assertIn(f'value="{tenant_id}" selected', page)

        catalogue = self.post("/playground", {
            "tenant_id": tenant_id,
            "action": "browse",
        }).get_data(as_text=True)
        with self.superuser() as cursor:
            cursor.execute(
                "select p.name, v.name, v.price, s.currency "
                "from public.products p "
                "join public.product_variants v "
                "on v.tenant_id = p.tenant_id and v.product_id = p.id "
                "join public.tenant_settings s on s.tenant_id = p.tenant_id "
                "where p.tenant_id = %s and p.is_active and v.is_active "
                "order by p.sort_order, p.name, v.sort_order, v.name limit 1",
                (tenant_id,),
            )
            active_catalogue_item = cursor.fetchone()
        if active_catalogue_item:
            product_name, variant_name, price, currency = active_catalogue_item
            self.assertIn(product_name, catalogue)
            self.assertIn(variant_name, catalogue)
            self.assertIn(f"{currency} {price:.2f}", catalogue)
        else:
            self.assertIn("No active products", catalogue)

    def test_playground_requires_an_authenticated_platform_administrator(self):
        anonymous = self.app.test_client()
        response = anonymous.get("/playground")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/login", response.headers["Location"])

    def test_playground_uses_the_selected_tenant_catalogue_and_configuration(self):
        self.onboard(SPECS[0])
        other_id = self.onboard(SPECS[1])

        response = self.post("/playground", {
            "tenant_id": other_id,
            "customer_message": "Can you repair my screen?",
        })
        page = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn("Grounded AI response", page)
        request = self.bot_provider.requests[-1]
        self.assertTrue(request.safe_test_mode)
        self.assertEqual(request.channel, "admin-playground")
        self.assertIn(SPECS[1]["product"]["name"], request.system_prompt)
        self.assertIn(SPECS[1]["persona"]["tone"], request.system_prompt)
        self.assertIn("warranty_days", request.system_prompt)
        self.assertNotIn(SPECS[0]["product"]["name"], request.system_prompt)

        catalogue = self.post("/playground", {
            "tenant_id": other_id,
            "action": "browse",
        }).get_data(as_text=True)
        self.assertIn(SPECS[1]["product"]["name"], catalogue)
        self.assertIn(SPECS[1]["variant"]["name"], catalogue)
        self.assertIn("USD 89.50", catalogue)
        self.assertNotIn(SPECS[0]["product"]["name"], catalogue)
        self.assertNotIn(SPECS[0]["variant"]["name"], catalogue)

    def test_conversation_history_is_scoped_to_the_selected_tenant(self):
        first = self.onboard(SPECS[0])
        second = self.onboard(SPECS[1])
        self.post("/playground", {"tenant_id": first, "customer_message": "alpha-only-message"})
        self.post("/playground", {"tenant_id": second, "customer_message": "beta-only-message"})

        first_page = self.page(f"/playground?tenant_id={first}")
        second_page = self.page(f"/playground?tenant_id={second}")
        self.assertIn("alpha-only-message", first_page)
        self.assertNotIn("beta-only-message", first_page)
        self.assertIn("beta-only-message", second_page)
        self.assertNotIn("alpha-only-message", second_page)

    def test_start_new_conversation_clears_only_the_selected_tenant(self):
        first = self.onboard(SPECS[0])
        second = self.onboard(SPECS[1])
        self.post("/playground", {"tenant_id": first, "customer_message": "clear-this"})
        self.post("/playground", {"tenant_id": second, "customer_message": "keep-this"})
        response = self.post("/playground", {"tenant_id": first, "start_new": "1"})
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("clear-this", self.page(f"/playground?tenant_id={first}"))
        self.assertIn("keep-this", self.page(f"/playground?tenant_id={second}"))

    def test_playground_rejects_bad_input_without_creating_live_side_effects(self):
        tenant_id = self.onboard(SPECS[0])
        with self.superuser() as cursor:
            cursor.execute(
                "select "
                "(select count(*) from public.audit_log where tenant_id = %s), "
                "(select count(*) from public.tenant_channels where tenant_id = %s), "
                "(select count(*) from public.subscriptions where tenant_id = %s)",
                (tenant_id, tenant_id, tenant_id),
            )
            before = cursor.fetchone()
        response = self.post("/playground", {
            "tenant_id": tenant_id,
            "customer_message": "x" * 5001,
        })
        page = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 400)
        self.assertIn("at most 2000", page)
        with self.superuser() as cursor:
            cursor.execute(
                "select "
                "(select count(*) from public.audit_log where tenant_id = %s), "
                "(select count(*) from public.tenant_channels where tenant_id = %s), "
                "(select count(*) from public.subscriptions where tenant_id = %s)",
                (tenant_id, tenant_id, tenant_id),
            )
            self.assertEqual(cursor.fetchone(), before)
        self.assertEqual(self.bot_provider.requests, [])

    def test_provider_errors_are_useful_and_do_not_leak_details(self):
        tenant_id = self.onboard(SPECS[0])
        self.bot_provider.failure = BotProviderError("The AI provider could not generate a response. Try again.")
        response = self.post(
            "/playground", {"tenant_id": tenant_id, "customer_message": "Hello"}
        )
        self.assertEqual(response.status_code, 502)
        page = response.get_data(as_text=True)
        self.assertIn("AI provider could not generate", page)
        self.assertNotIn("OPENAI_API_KEY", page)

    def test_successful_playground_response_has_no_live_side_effects(self):
        tenant_id = self.onboard(SPECS[0])
        with self.superuser() as cursor:
            cursor.execute(
                "select "
                "(select count(*) from public.audit_log where tenant_id = %s), "
                "(select count(*) from public.tenant_channels where tenant_id = %s), "
                "(select count(*) from public.subscriptions where tenant_id = %s)",
                (tenant_id, tenant_id, tenant_id),
            )
            before = cursor.fetchone()
        response = self.post(
            "/playground", {"tenant_id": tenant_id, "customer_message": "Book this now"}
        )
        self.assertEqual(response.status_code, 200)
        with self.superuser() as cursor:
            cursor.execute(
                "select "
                "(select count(*) from public.audit_log where tenant_id = %s), "
                "(select count(*) from public.tenant_channels where tenant_id = %s), "
                "(select count(*) from public.subscriptions where tenant_id = %s)",
                (tenant_id, tenant_id, tenant_id),
            )
            self.assertEqual(cursor.fetchone(), before)


class MemberAndSubscriptionTests(DashboardTestCase):
    def setUp(self):
        super().setUp()
        self.login()
        self.tenant = self.create_tenant(SPECS[0])

    def test_members_require_a_valid_uuid_and_role(self):
        base = f"/tenants/{self.tenant}/members"
        self.post(base, {"user_id": "nope", "role": "owner"})
        self.post(base, {"user_id": str(uuid.uuid4()), "role": "admin"})
        with self.superuser() as cursor:
            cursor.execute("select count(*) from public.tenant_users where tenant_id = %s", (self.tenant,))
            self.assertEqual(cursor.fetchone()[0], 0)
        self.post(base, {"user_id": str(uuid.uuid4()), "role": "owner"})
        self.assertIn("member.added", self.audit_actions(self.tenant))

    def test_subscription_can_be_set_and_changed(self):
        base = f"/tenants/{self.tenant}/subscription"
        with self.superuser() as cursor:
            cursor.execute("select id::text from public.plans where is_active limit 1")
            plan_id = cursor.fetchone()[0]
        self.post(base, {"plan_id": plan_id, "status": "trial"})
        self.post(base, {"plan_id": plan_id, "status": "active"})
        with self.superuser() as cursor:
            cursor.execute("select status from public.subscriptions where tenant_id = %s", (self.tenant,))
            self.assertEqual(cursor.fetchall(), [("active",)])
        self.post(base, {"plan_id": plan_id, "status": "bogus"})
        self.assertEqual(self.audit_actions(self.tenant).count("subscription.updated"), 2)


if __name__ == "__main__":
    unittest.main()
