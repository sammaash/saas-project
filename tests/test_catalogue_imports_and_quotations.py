"""Database-free contracts for Excel catalogue imports and quotation rendering."""

from __future__ import annotations

import io
import unittest
import uuid
from decimal import Decimal

from openpyxl import load_workbook

from platform_admin.services import catalogue_import, playground, quotations
from platform_admin.services.validation import ServiceError


def _saved(workbook) -> bytes:
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class CatalogueTemplateTests(unittest.TestCase):
    def test_template_has_instructions_headers_examples_and_validations(self):
        workbook = load_workbook(io.BytesIO(catalogue_import.template_bytes()))
        self.assertEqual(workbook.sheetnames, ["Instructions", "Products"])
        products = workbook["Products"]
        self.assertEqual(
            tuple(cell.value for cell in products[1]), catalogue_import.HEADERS
        )
        self.assertEqual(products["A2"].value, "CAKE-001-1KG")
        validations = list(products.data_validations.dataValidation)
        self.assertTrue(any("IN_STOCK" in str(item.formula1) for item in validations))
        self.assertIn("SKU is unique", "\n".join(
            str(cell.value or "") for cell in workbook["Instructions"]["A"]
        ))

    def test_parser_rejects_wrong_type_headers_duplicates_and_malformed_rows(self):
        with self.assertRaisesRegex(ServiceError, r"\.xlsx"):
            catalogue_import.parse_workbook("catalogue.csv", b"not excel")

        workbook = load_workbook(io.BytesIO(catalogue_import.template_bytes()))
        workbook["Products"]["A1"] = "Wrong"
        with self.assertRaisesRegex(ServiceError, "headers"):
            catalogue_import.parse_workbook("catalogue.xlsx", _saved(workbook))

        workbook = load_workbook(io.BytesIO(catalogue_import.template_bytes()))
        workbook["Products"]["A3"] = workbook["Products"]["A2"].value
        rows, errors = catalogue_import.parse_workbook(
            "catalogue.xlsx", _saved(workbook)
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(len(errors), 1)
        self.assertIn("duplicate SKU", errors[0]["message"])

        workbook = load_workbook(io.BytesIO(catalogue_import.template_bytes()))
        workbook["Products"]["F2"] = "not-a-price"
        rows, errors = catalogue_import.parse_workbook(
            "catalogue.xlsx", _saved(workbook)
        )
        self.assertEqual([row["row_number"] for row in rows], [3])
        self.assertIn("Price must be a number", errors[0]["message"])

    def test_parser_preserves_blank_optional_cells(self):
        workbook = load_workbook(io.BytesIO(catalogue_import.template_bytes()))
        products = workbook["Products"]
        products.delete_rows(3)
        for column in (3, 4, 5, 9, 10, 11):
            products.cell(2, column).value = None
        rows, errors = catalogue_import.parse_workbook(
            "catalogue.xlsx", _saved(workbook)
        )
        self.assertEqual(errors, [])
        self.assertIsNone(rows[0]["variant"])
        self.assertIsNone(rows[0]["stock_quantity"])
        self.assertIsNone(rows[0]["image_url"])

    def test_parser_reports_negative_quantities_as_row_errors(self):
        workbook = load_workbook(io.BytesIO(catalogue_import.template_bytes()))
        workbook["Products"].delete_rows(3)
        workbook["Products"]["I2"] = -1
        rows, errors = catalogue_import.parse_workbook(
            "catalogue.xlsx", _saved(workbook)
        )
        self.assertEqual(rows, [])
        self.assertEqual(errors[0]["row_number"], 2)
        self.assertIn("Stock Quantity must be between 0", errors[0]["message"])


class AvailabilityTests(unittest.TestCase):
    def setUp(self):
        self.product_id = str(uuid.uuid4())
        self.variant_id = str(uuid.uuid4())

    def _catalogue(self, availability="IN_STOCK", stock_quantity=None):
        return [
            {
                "id": self.product_id,
                "name": "Cake",
                "variants": [
                    {
                        "id": self.variant_id,
                        "name": "Large",
                        "price": 10,
                        "availability": availability,
                        "stock_quantity": stock_quantity,
                    }
                ],
            }
        ]

    def test_out_of_stock_cannot_be_added_but_can_be_removed_from_stale_cart(self):
        catalogue = self._catalogue("OUT_OF_STOCK")
        with self.assertRaisesRegex(ServiceError, "out of stock"):
            playground.add_cart_item(
                catalogue, [], self.product_id, self.variant_id, 1
            )
        stale = [
            {
                "product_id": self.product_id,
                "variant_id": self.variant_id,
                "quantity": 1,
            }
        ]
        self.assertEqual(
            playground.change_cart_item(
                catalogue, stale, self.product_id, self.variant_id, 0
            ),
            [],
        )

    def test_configured_stock_caps_cart_quantity(self):
        with self.assertRaisesRegex(ServiceError, "Only 2"):
            playground.add_cart_item(
                self._catalogue(stock_quantity=2),
                [],
                self.product_id,
                self.variant_id,
                3,
            )

    def test_made_to_order_is_allowed_and_labelled(self):
        catalogue = self._catalogue("MADE_TO_ORDER")
        cart = playground.add_cart_item(
            catalogue, [], self.product_id, self.variant_id, 2
        )
        summary = playground.cart_summary(catalogue, cart, "USD")
        self.assertEqual(summary["items"][0]["availability"], "MADE_TO_ORDER")
        self.assertIn("Made to order", summary["items"][0]["availability_note"])


class QuotationTests(unittest.TestCase):
    def test_absent_tax_configuration_does_not_invent_a_rate(self):
        rate, treatment = quotations._tax_configuration({})
        self.assertIsNone(rate)
        self.assertIn("not configured", treatment)

    def test_pdf_contains_a_valid_pdf_document(self):
        quotation = {
            "business_name": "Example Bakery",
            "quotation_number": "Q-20250101-ABC12345",
            "customer_name": "Sample Customer",
            "customer_contact": "customer@example.test",
            "issue_date": "2025-01-01",
            "expiry_date": "2025-01-15",
            "currency": "USD",
            "subtotal": Decimal("20.00"),
            "tax_amount": None,
            "total": Decimal("20.00"),
            "tax_treatment": "Tax is not configured for this tenant.",
            "terms": None,
            "lines": [
                {
                    "sku": "CAKE-1",
                    "description": "Cake - Large",
                    "availability": "MADE_TO_ORDER",
                    "quantity": 2,
                    "unit_price": Decimal("10.00"),
                    "subtotal": Decimal("20.00"),
                }
            ],
        }
        document = quotations.quotation_pdf(quotation)
        self.assertTrue(document.startswith(b"%PDF-"))
        self.assertGreater(len(document), 1000)


if __name__ == "__main__":
    unittest.main()
