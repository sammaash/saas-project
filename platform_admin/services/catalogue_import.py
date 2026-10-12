"""Excel catalogue templates, validation previews and atomic SKU upserts."""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any
from urllib.parse import urlparse

from saas.tenant_context import tenant_transaction

from ..config import Settings
from . import audit, validation as v

HEADERS = (
    "SKU",
    "Product Name",
    "Variant",
    "Category",
    "Description",
    "Price",
    "Currency",
    "Availability",
    "Stock Quantity",
    "Reorder Level",
    "Image URL",
)
AVAILABILITY_VALUES = frozenset({"IN_STOCK", "OUT_OF_STOCK", "MADE_TO_ORDER"})
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 5000


def template_bytes() -> bytes:
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Font, PatternFill  # noqa: PLC0415
    from openpyxl.worksheet.datavalidation import DataValidation  # noqa: PLC0415

    workbook = Workbook()
    products = workbook.active
    products.title = "Products"
    products.append(HEADERS)
    products.append(
        (
            "CAKE-001-1KG", "Lemon Cake", "1 kg", "Cakes", "Fresh lemon sponge",
            35.00, "USD", "IN_STOCK", 12, 4, "https://example.com/lemon-cake.jpg",
        )
    )
    products.append(
        (
            "CAKE-001-2KG", "Lemon Cake", "2 kg", "Cakes", "Fresh lemon sponge",
            60.00, "USD", "MADE_TO_ORDER", None, None, None,
        )
    )
    products.freeze_panes = "A2"
    products.auto_filter.ref = "A1:K3"
    for cell in products[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    widths = (20, 24, 18, 18, 36, 14, 12, 20, 16, 16, 42)
    for index, width in enumerate(widths, 1):
        products.column_dimensions[products.cell(1, index).column_letter].width = width
    availability = DataValidation(
        type="list",
        formula1='"IN_STOCK,OUT_OF_STOCK,MADE_TO_ORDER"',
        allow_blank=False,
    )
    products.add_data_validation(availability)
    availability.add("H2:H5001")
    price = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
    products.add_data_validation(price)
    price.add("F2:F5001")
    quantity = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0")
    products.add_data_validation(quantity)
    quantity.add("I2:J5001")

    instructions = workbook.create_sheet("Instructions", 0)
    instructions.append(("Catalogue import instructions",))
    instructions["A1"].font = Font(bold=True, size=14)
    notes = (
        "Use the Products worksheet and keep the header row unchanged.",
        "SKU, Product Name, Price, Currency and Availability are required.",
        "SKU is unique within a tenant and is used to update an existing variant.",
        "Availability must be IN_STOCK, OUT_OF_STOCK or MADE_TO_ORDER.",
        "Blank optional cells do not erase existing values during an update.",
        "Stock Quantity and Reorder Level are optional non-negative whole numbers.",
        "Currency must be a three-letter ISO code matching the tenant currency.",
        "Review the preview before applying. Omitted rows are never deleted or disabled.",
        "The example rows in Products are illustrative; replace or remove them before import.",
    )
    for note in notes:
        instructions.append((note,))
    instructions.column_dimensions["A"].width = 110

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_text(value: Any, label: str, row_number: int, *, required: bool, limit: int) -> str | None:
    try:
        return v.text(value, label, required=required, max_length=limit)
    except v.ServiceError as error:
        raise v.ServiceError(f"Row {row_number}: {error.message}") from error


def _optional_integer(value: Any, label: str, row_number: int) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return v.integer(value, label)
    except v.ServiceError as error:
        raise v.ServiceError(f"Row {row_number}: {error.message}") from error


def _parse_row(values: tuple[Any, ...], row_number: int) -> dict[str, Any]:
    padded = (*values, *(None for _ in range(len(HEADERS) - len(values))))
    sku = _cell_text(padded[0], "SKU", row_number, required=True, limit=100)
    product_name = _cell_text(
        padded[1], "Product Name", row_number, required=True, limit=200
    )
    variant = _cell_text(padded[2], "Variant", row_number, required=False, limit=200)
    category = _cell_text(padded[3], "Category", row_number, required=False, limit=100)
    description = _cell_text(
        padded[4], "Description", row_number, required=False, limit=2000
    )
    try:
        price = v.price(padded[5])
        currency = v.currency(padded[6])
    except v.ServiceError as error:
        raise v.ServiceError(f"Row {row_number}: {error.message}") from error
    availability = str(padded[7] or "").strip().upper()
    if availability not in AVAILABILITY_VALUES:
        raise v.ServiceError(
            f"Row {row_number}: Availability must be IN_STOCK, OUT_OF_STOCK or "
            "MADE_TO_ORDER."
        )
    stock_quantity = _optional_integer(padded[8], "Stock Quantity", row_number)
    reorder_level = _optional_integer(padded[9], "Reorder Level", row_number)
    image_url = _cell_text(padded[10], "Image URL", row_number, required=False, limit=2000)
    if image_url:
        parsed = urlparse(image_url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            raise v.ServiceError(f"Row {row_number}: Image URL must be an HTTP(S) URL.")
    return {
        "row_number": row_number,
        "sku": sku,
        "product_name": product_name,
        "variant": variant,
        "category": category,
        "description": description,
        "price": f"{price:.2f}",
        "currency": currency,
        "availability": availability,
        "stock_quantity": stock_quantity,
        "reorder_level": reorder_level,
        "image_url": image_url,
    }


def parse_workbook(filename: str, content: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if not filename.lower().endswith(".xlsx"):
        raise v.ServiceError("Catalogue imports must use an .xlsx file.")
    if not content:
        raise v.ServiceError("The uploaded workbook is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise v.ServiceError("The workbook must be 2 MB or smaller.")
    try:
        from openpyxl import load_workbook  # noqa: PLC0415

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise v.ServiceError("The uploaded file is not a valid .xlsx workbook.") from error
    if "Products" not in workbook.sheetnames:
        raise v.ServiceError("The workbook must contain a Products worksheet.")
    sheet = workbook["Products"]
    header = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    if header != HEADERS:
        raise v.ServiceError("The Products worksheet headers do not match the template.")

    rows: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen_skus: set[str] = set()
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True), 2
    ):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        if len(rows) + len(errors) >= MAX_ROWS:
            raise v.ServiceError(f"A workbook may contain at most {MAX_ROWS} data rows.")
        try:
            parsed = _parse_row(values, row_number)
            normalized_sku = parsed["sku"].casefold()
            if normalized_sku in seen_skus:
                raise v.ServiceError(f"Row {row_number}: duplicate SKU {parsed['sku']!r}.")
            seen_skus.add(normalized_sku)
            rows.append(parsed)
        except v.ServiceError as error:
            errors.append({"row_number": row_number, "message": error.message})
    if not rows and not errors:
        raise v.ServiceError("The Products worksheet contains no catalogue rows.")
    return rows, errors


def _row_dict(cursor: Any, row: tuple[Any, ...]) -> dict[str, Any]:
    return dict(zip((column.name for column in cursor.description), row, strict=True))


def preview_import(
    settings: Settings,
    admin_id: str,
    tenant_id: str,
    filename: str,
    content: bytes,
) -> dict[str, Any]:
    try:
        rows, errors = parse_workbook(filename, content)
    except v.ServiceError:
        audit.record(
            settings,
            admin_id,
            tenant_id,
            "catalogue.import_rejected",
            "catalogue_import",
            metadata={"error_count": 1, "stage": "workbook_validation"},
        )
        raise
    with v.translated(), tenant_transaction(tenant_id, settings.app_database_url) as cursor:
        cursor.execute(
            "select currency, catalogue_version from public.tenant_settings "
            "where tenant_id = app.current_tenant_id()"
        )
        settings_row = cursor.fetchone()
        if settings_row is None:
            raise v.ServiceError("Tenant settings were not found.", 404)
        currency, version = settings_row
        for row in rows:
            if row["currency"] != currency:
                errors.append(
                    {
                        "row_number": row["row_number"],
                        "message": (
                            f"Row {row['row_number']}: currency {row['currency']} does not "
                            f"match tenant currency {currency}."
                        ),
                    }
                )
        cursor.execute(
            """
            select v.id, v.sku, v.name as variant_name, v.price, v.availability,
                   v.stock_quantity, v.reorder_level, p.name as product_name,
                   p.category, p.description, p.attributes
            from public.product_variants v
            join public.products p on p.id = v.product_id and p.tenant_id = v.tenant_id
            where v.tenant_id = app.current_tenant_id()
            """
        )
        existing = {
            str(item["sku"]).casefold(): item
            for item in (_row_dict(cursor, value) for value in cursor.fetchall())
        }
        counts = {"new": 0, "updates": 0, "unchanged": 0, "availability_changes": 0}
        for row in rows:
            current = existing.get(row["sku"].casefold())
            if current is None:
                row["change"] = "new"
                row["availability_change"] = False
                counts["new"] += 1
                continue
            availability_change = current["availability"] != row["availability"]
            changed = (
                current["product_name"] != row["product_name"]
                or (row["variant"] is not None and current["variant_name"] != row["variant"])
                or Decimal(current["price"]) != Decimal(row["price"])
                or availability_change
                or (
                    row["stock_quantity"] is not None
                    and current["stock_quantity"] != row["stock_quantity"]
                )
                or (
                    row["reorder_level"] is not None
                    and current["reorder_level"] != row["reorder_level"]
                )
                or (row["category"] is not None and current["category"] != row["category"])
                or (
                    row["description"] is not None
                    and current["description"] != row["description"]
                )
                or (
                    row["image_url"] is not None
                    and (current["attributes"] or {}).get("image_url") != row["image_url"]
                )
            )
            row["change"] = "update" if changed else "unchanged"
            row["availability_change"] = availability_change
            counts["updates" if changed else "unchanged"] += 1
            if availability_change:
                counts["availability_changes"] += 1

        batch_id = None
        if not errors:
            from psycopg.types.json import Jsonb  # noqa: PLC0415

            batch_id = str(uuid.uuid4())
            cursor.execute(
                """
                insert into public.catalogue_import_batches
                    (id, tenant_id, actor_user_id, expected_version, rows, counts)
                values (%s, app.current_tenant_id(), %s, %s, %s, %s)
                """,
                (batch_id, admin_id, version, Jsonb(rows), Jsonb(counts)),
            )
            audit.record(
                settings,
                admin_id,
                tenant_id,
                "catalogue.import_previewed",
                "catalogue_import",
                batch_id,
                counts,
            )
        else:
            audit.record(
                settings,
                admin_id,
                tenant_id,
                "catalogue.import_rejected",
                "catalogue_import",
                metadata={"error_count": len(errors), **counts},
            )
    return {"batch_id": batch_id, "rows": rows, "errors": errors, "counts": counts}


def _upsert_row(cursor: Any, row: dict[str, Any]) -> None:
    from psycopg.types.json import Jsonb  # noqa: PLC0415

    cursor.execute(
        """
        select v.id
        from public.product_variants v
        where v.tenant_id = app.current_tenant_id() and lower(btrim(v.sku)) = lower(btrim(%s))
        """,
        (row["sku"],),
    )
    existing = cursor.fetchone()
    cursor.execute(
        "select id, attributes from public.products "
        "where tenant_id = app.current_tenant_id() and name_normalized = lower(btrim(%s))",
        (row["product_name"],),
    )
    product = cursor.fetchone()
    if product:
        product_id, current_attributes = product
        attributes = dict(current_attributes or {})
    else:
        attributes = {}
        if row["image_url"] is not None:
            attributes["image_url"] = row["image_url"]
        cursor.execute(
            """
            insert into public.products
                (tenant_id, name, description, category, attributes, is_active)
            values (app.current_tenant_id(), %s, %s, %s, %s, true)
            returning id
            """,
            (
                row["product_name"],
                row["description"],
                row["category"],
                Jsonb(attributes),
            ),
        )
        product_id = cursor.fetchone()[0]
    if row["image_url"] is not None:
        attributes["image_url"] = row["image_url"]
    product_sets = ["name = %s", "attributes = %s"]
    product_values: list[Any] = [row["product_name"], Jsonb(attributes)]
    for field in ("description", "category"):
        if row[field] is not None:
            product_sets.append(f"{field} = %s")
            product_values.append(row[field])
    cursor.execute(
        f"update public.products set {', '.join(product_sets)} "
        "where id = %s and tenant_id = app.current_tenant_id()",
        (*product_values, product_id),
    )
    if existing:
        variant_sets = [
            "product_id = %s", "price = %s", "availability = %s", "is_active = true"
        ]
        variant_values: list[Any] = [
            product_id, Decimal(row["price"]), row["availability"]
        ]
        if row["variant"] is not None:
            variant_sets.append("name = %s")
            variant_values.append(row["variant"])
        for field in ("stock_quantity", "reorder_level"):
            if row[field] is not None:
                variant_sets.append(f"{field} = %s")
                variant_values.append(row[field])
        cursor.execute(
            f"update public.product_variants set {', '.join(variant_sets)} "
            "where id = %s and tenant_id = app.current_tenant_id()",
            (*variant_values, existing[0]),
        )
    else:
        cursor.execute(
            """
            insert into public.product_variants
                (tenant_id, product_id, code, sku, name, price, availability,
                 stock_quantity, reorder_level, is_active)
            values (app.current_tenant_id(), %s, %s, %s, %s, %s, %s, %s, %s, true)
            """,
            (
                product_id,
                row["sku"],
                row["sku"],
                row["variant"] or row["product_name"],
                Decimal(row["price"]),
                row["availability"],
                row["stock_quantity"],
                row["reorder_level"],
            ),
        )


def apply_import(
    settings: Settings, admin_id: str, tenant_id: str, batch_id: str
) -> dict[str, int]:
    batch_id = v.uuid_value(batch_id, "Import preview")
    now = datetime.now(timezone.utc)
    with v.translated(), tenant_transaction(tenant_id, settings.app_database_url) as cursor:
        cursor.execute(
            """
            select actor_user_id, status, expected_version, rows, counts, expires_at
            from public.catalogue_import_batches
            where id = %s and tenant_id = app.current_tenant_id()
            for update
            """,
            (batch_id,),
        )
        batch = cursor.fetchone()
        if batch is None:
            raise v.ServiceError("Import preview not found.", 404)
        actor_user_id, status, expected_version, rows, counts, expires_at = batch
        if str(actor_user_id) != str(admin_id):
            raise v.ServiceError("Only the administrator who created this preview can apply it.", 403)
        if status != "preview":
            raise v.ServiceError("This import preview has already been used.", 409)
        if expires_at <= now:
            cursor.execute(
                "update public.catalogue_import_batches set status = 'expired' where id = %s",
                (batch_id,),
            )
            raise v.ServiceError("This import preview has expired. Upload the workbook again.", 409)
        cursor.execute(
            "select catalogue_version from public.tenant_settings "
            "where tenant_id = app.current_tenant_id() for update"
        )
        current_version = cursor.fetchone()[0]
        if current_version != expected_version:
            raise v.ServiceError(
                "The catalogue changed after this preview. Upload the workbook again.", 409
            )
        for row in rows:
            _upsert_row(cursor, row)
        cursor.execute(
            "update public.catalogue_import_batches "
            "set status = 'applied', applied_at = now() where id = %s",
            (batch_id,),
        )
        audit.record(
            settings,
            admin_id,
            tenant_id,
            "catalogue.import_applied",
            "catalogue_import",
            batch_id,
            dict(counts),
        )
    return dict(counts)
